#!/usr/bin/env python3
"""
Выгрузка и загрузка файлов: .txt и .docx, имена по-русски.

Зачем. Три пути жили без тестов и были сломаны (AUDIT_UNIFIED.md, F1, F2):
- имя проекта шло в Content-Disposition как есть; заголовки HTTP
  кодируются в latin-1, и на реальном сервере Werkzeug выгрузка любого
  проекта с русским именем обрывалась UnicodeEncodeError. Тестовый клиент
  Flask заголовки не кодирует, поэтому здесь это проверяется явно —
  тем же .encode("latin-1"), что делает http.server;
- python-docx не было в зависимостях: экспорт и импорт .docx отвечали
  «установи python-docx»;
- загрузка главы декодировала любой файл как UTF-8, и .docx (zip-архив)
  молча сохранялся в главу мусором.
"""

import io
import sys
from unittest.mock import MagicMock
from urllib.parse import unquote

import pytest

for _mod in ("openai", "anthropic"):
    sys.modules.setdefault(_mod, MagicMock())
sys.modules["openai"].OpenAI = MagicMock()

NAME = "Тёмная «серия» №1"


def _filename_star(header: str) -> str:
    """Имя из filename*=UTF-8''… — то, что увидит браузер."""
    part = header.split("filename*=", 1)[1]
    assert part.startswith("UTF-8''"), header
    return unquote(part[len("UTF-8''"):])


def _assert_sendable(header: str) -> None:
    header.encode("latin-1")   # ровно так заголовок пишет http.server


@pytest.fixture
def client():
    import logging
    from web.app import app
    app.config["TESTING"] = True
    app.config["SECRET_KEY"] = "test-secret"
    logging.disable(logging.CRITICAL)
    yield app.test_client()
    logging.disable(logging.NOTSET)


@pytest.fixture
def project():
    from engine.db import create_project, save_chapter, set_active_project
    pid = create_project(NAME, "детектив")
    set_active_project(pid)
    save_chapter(pid, 1, "Первый абзац главы.\nВторой абзац главы.", "Начало")
    return pid


def _docx_bytes(*paragraphs: str) -> bytes:
    import docx
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ─── Заголовок ───────────────────────────────────────────────────────────────

class TestAttachmentHeader:
    @pytest.mark.parametrize("filename", [
        "Моя книга.txt", f"{NAME}.docx", "plain.txt", "без_расширения",
        'кавычки " и ; точка с запятой.txt',
    ])
    def test_sendable_and_exact(self, filename):
        from web.blueprints.helpers import attachment_header
        header = attachment_header(filename)
        _assert_sendable(header)
        assert _filename_star(header) == filename

    def test_ascii_fallback_keeps_extension(self):
        from web.blueprints.helpers import attachment_header
        header = attachment_header("Моя книга.docx")
        assert 'filename="export.docx"' in header

    def test_ascii_name_kept(self):
        from web.blueprints.helpers import attachment_header
        assert 'filename="chapter_3.txt"' in attachment_header("chapter_3.txt")


# ─── Выгрузка ────────────────────────────────────────────────────────────────

class TestExport:
    def test_txt_with_russian_name(self, client, project):
        r = client.get("/export/txt")
        assert r.status_code == 200
        header = r.headers["Content-Disposition"]
        _assert_sendable(header)
        assert _filename_star(header) == f"{NAME}.txt"
        assert "Второй абзац главы." in r.get_data(as_text=True)

    def test_chapter_txt(self, client, project):
        r = client.get("/export/chapter/1/txt")
        assert r.status_code == 200
        _assert_sendable(r.headers["Content-Disposition"])

    def test_docx_with_russian_name(self, client, project):
        import docx
        r = client.get("/export/docx")
        assert r.status_code == 200, "экспорт .docx не отдал файл"
        assert r.mimetype.endswith("wordprocessingml.document")
        header = r.headers["Content-Disposition"]
        _assert_sendable(header)
        assert _filename_star(header) == f"{NAME}.docx"
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.data)).paragraphs)
        assert NAME in text
        assert "Глава 1: Начало" in text
        assert "Второй абзац главы." in text


# ─── Загрузка ────────────────────────────────────────────────────────────────

class TestUpload:
    def _upload(self, client, filename, payload, number=2):
        return client.post("/chapter/upload", data={
            "number": str(number), "title": "",
            "file": (io.BytesIO(payload), filename),
        }, content_type="multipart/form-data")

    def test_chapter_from_docx_is_text_not_zip(self, client, project):
        from engine.db import get_chapter
        self._upload(client, "глава.docx", _docx_bytes("Строка один.", "Строка два."))
        content = get_chapter(project, 2)["content"]
        assert content == "Строка один.\nСтрока два."

    def test_chapter_from_txt(self, client, project):
        from engine.db import get_chapter
        self._upload(client, "глава.txt", "Текст главы.".encode("utf-8"))
        assert get_chapter(project, 2)["content"] == "Текст главы."

    def test_unsupported_format_is_not_saved(self, client, project):
        from engine.db import get_chapter
        self._upload(client, "глава.pdf", b"%PDF-1.4 binary")
        assert get_chapter(project, 2) is None

    def test_state_import_reads_docx(self, client, project, monkeypatch):
        import engine.pipeline as pipeline
        seen = {}

        def fake_call_json(model, system, prompt, max_tokens=0):
            seen["prompt"] = prompt
            return {"global_state": "g", "plot_matrix": "p", "memory_graph": "m"}

        monkeypatch.setattr(pipeline, "call_json", fake_call_json)
        r = client.post("/state/import", data={
            "model": "anthropic::claude-test",
            "file": (io.BytesIO(_docx_bytes("Аня — сыщица.")), "материал.docx"),
        }, content_type="multipart/form-data")
        assert r.status_code == 200, r.get_json()
        assert "Аня — сыщица." in seen["prompt"]

    def test_state_import_rejects_unsupported(self, client, project):
        r = client.post("/state/import", data={
            "model": "anthropic::claude-test",
            "file": (io.BytesIO(b"x"), "материал.pdf"),
        }, content_type="multipart/form-data")
        assert r.status_code == 400
