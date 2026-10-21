"""Blueprint: главы, экспорт, L3."""
from flask import Blueprint, render_template, request, jsonify, redirect, url_for, flash, Response
from engine.db import (delete_chapter, get_chapter, get_chapters, rollback_chapter, rollback_info,
                       save_chapter)
from .helpers import (NO_MODEL, attachment_header, get_current_project, get_cheap_model,
                      page_model, read_uploaded_text)

bp = Blueprint("chapters", __name__)

# Главная страница объявлена прямо на app (web/app.py: @app.route("/")),
# поэтому её endpoint — "index", без префикса блупринта. Раньше здесь
# стояло url_for("main.index") — имя из планировщика, где такой блупринт
# есть. В Fiction Engine его нет, и любой редирект «нет проекта → на
# главную» падал с BuildError, то есть 500 вместо страницы выбора проекта.


@bp.route("/chapter/upload", methods=["POST"])
def chapter_upload():
    current = get_current_project()
    if not current:
        flash("Сначала выбери проект", "error")
        return redirect(url_for("index"))
    number = request.form.get("number", type=int)
    title  = request.form.get("title", "").strip()
    file   = request.files.get("file")
    text   = request.form.get("text", "").strip()
    if file and file.filename:
        try:
            content = read_uploaded_text(file)
        except ValueError as e:
            flash(str(e), "error")
            return redirect(url_for("index"))
    elif text:
        content = text
    else:
        flash("Нужен файл или текст", "error")
        return redirect(url_for("index"))
    if not number:
        flash("Укажи номер главы", "error")
        return redirect(url_for("index"))
    replaced = bool(get_chapter(current["id"], number))
    save_chapter(current["id"], number, content, title)
    flash(f"Глава {number} сохранена ({len(content.split())} слов)" + (
        ". L3, анализ, State и библия — ещё от прежнего текста: «Пересчитать по тексту» на странице главы"
        if replaced else ""), "success")
    return redirect(url_for("index"))


@bp.route("/chapter/<int:num>/view")
def chapter_view(num):
    current = get_current_project()
    if not current:
        return redirect(url_for("index"))
    ch = get_chapter(current["id"], num)
    return render_template("chapter.html", chapter=ch, current=current)


@bp.route("/chapter/<int:num>/save", methods=["POST"])
def chapter_save(num):
    """
    Правка текста главы автором прямо на странице главы. Раньше главу можно
    было только перезаписать целиком через «Обзор → Текст» с копированием.
    Сразу ничего не пересчитывается: правка обычно точечная, а пересчёт — вызовы
    моделей. Ответ stale — страница предлагает «Пересчитать по тексту»
    (/chapter/N/recompute): без него L3, анализ, State и библия остаются от
    прежнего текста и уходят писателю следующих глав (разбор 21.10).
    """
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    ch = get_chapter(current["id"], num)
    if not ch:
        return jsonify({"error": "Глава не найдена"}), 404
    text = ((request.json or {}).get("text") or "").strip()
    if len(text) < 100:
        return jsonify({"error": "Текст главы пуст или слишком короткий"}), 400
    save_chapter(current["id"], num, text, ch.get("title") or f"Глава {num}")
    return jsonify({"ok": True, "word_count": len(text.split()),
                    "stale": text != (ch.get("content") or "").strip()})


@bp.route("/chapter/<int:num>/recompute", methods=["POST"])
def chapter_recompute(num):
    """
    Пересчитать производное главы по её нынешнему тексту — тем же путём, что
    «Принять»: убрать посчитанное по прежнему тексту (prepare_chapter_replace),
    затем в фоне State, анализ главы, L3, заметка, оценка.
    """
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    ch = get_chapter(current["id"], num)
    if not ch:
        return jsonify({"error": "Глава не найдена"}), 404
    from engine.db import chapter_writers
    from engine.pipeline_config import default_role_models
    model = chapter_writers(current["id"]).get(num) or default_role_models().get("gen", "")
    model = page_model(model)
    if not model:
        return jsonify({"error": NO_MODEL}), 400
    from engine.db_chapters import prepare_chapter_replace
    from .generate_pipeline import _start_state_update
    prepare_chapter_replace(current["id"], num)
    return jsonify({"ok": True, "state_job": _start_state_update(current["id"], num, model, ch["content"])})


@bp.route("/chapter/<int:num>/delete", methods=["POST"])
def chapter_delete(num):
    """Удалить главу. Удалить её раньше было нельзя ни из одного места."""
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    if not delete_chapter(current["id"], num):
        return jsonify({"error": "Глава не найдена"}), 404
    return jsonify({"ok": True})


@bp.route("/chapter/<int:num>/rollback")
def chapter_rollback_info(num):
    """Можно ли откатить главу целиком и к какому снимку State."""
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    if not get_chapter(current["id"], num):
        return jsonify({"error": "Глава не найдена"}), 404
    info = rollback_info(current["id"], num)
    snap = info["snapshot"]
    return jsonify({"last": info["last"], "snapshot": snap and {"id": snap["id"], "saved_at": snap["saved_at"]}})


@bp.route("/chapter/<int:num>/rollback", methods=["POST"])
def chapter_rollback(num):
    """Откатить последнюю главу: State до её анализа, глава и производное от неё."""
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    if not get_chapter(current["id"], num):
        return jsonify({"error": "Глава не найдена"}), 404
    r = rollback_chapter(current["id"], num)
    if not r["ok"]:
        return jsonify({"error": r.get("error") or "Глава не откатилась"}), 400
    return jsonify({"ok": True, "state_restored": r["state_restored"],
                    "promises_reopened": r["promises_reopened"]})


@bp.route("/export/txt")
def export_txt():
    current = get_current_project()
    if not current:
        return redirect(url_for("index"))
    chapters = get_chapters(current["id"])
    lines = []
    for ch in chapters:
        full = get_chapter(current["id"], ch["number"])
        lines.append(f"=== Глава {ch['number']} ===\n\n{full['content']}\n\n")
    return Response("\n".join(lines), mimetype="text/plain",
                    headers={"Content-Disposition": attachment_header(f"{current['name']}.txt")})


@bp.route("/export/chapter/<int:num>/txt")
def export_chapter_txt(num):
    current = get_current_project()
    if not current:
        return redirect(url_for("index"))
    ch = get_chapter(current["id"], num)
    if not ch:
        flash("Глава не найдена", "error")
        return redirect(url_for("index"))
    return Response(ch["content"], mimetype="text/plain",
                    headers={"Content-Disposition": attachment_header(f"chapter_{num}.txt")})




@bp.route("/export/docx")
def export_docx():
    """Экспорт всех глав в один DOCX файл."""
    current = get_current_project()
    if not current:
        return redirect(url_for("index"))
    chapters = get_chapters(current["id"])
    if not chapters:
        flash("Нет глав для экспорта", "error")
        return redirect(url_for("index"))

    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import io

        doc = Document()

        # Стиль документа
        style = doc.styles["Normal"]
        style.font.name = "Georgia"
        style.font.size = Pt(12)

        # Заголовок книги
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_p.add_run(current["name"])
        run.font.size = Pt(24)
        run.bold = True

        if current.get("genre"):
            sub_p = doc.add_paragraph()
            sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub_run = sub_p.add_run(current["genre"])
            sub_run.font.size = Pt(12)
            sub_run.italic = True

        doc.add_page_break()

        for ch in chapters:
            full = get_chapter(current["id"], ch["number"])
            if not full or not full.get("content"):
                continue

            # Заголовок главы
            h = doc.add_paragraph()
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hr = h.add_run(f"Глава {ch['number']}" +
                          (f": {ch['title']}" if ch.get("title") else ""))
            hr.font.size = Pt(16)
            hr.bold = True

            doc.add_paragraph()  # отступ

            # Текст главы — разбиваем по абзацам
            for para_text in full["content"].split("\n"):
                p = doc.add_paragraph(para_text)
                p.paragraph_format.first_line_indent = Inches(0.3)

            doc.add_page_break()

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        return Response(
            buf.read(),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": attachment_header(f"{current['name']}.docx")}
        )
    except ImportError:
        flash("Установи python-docx: pip install python-docx", "error")
        return redirect(url_for("index"))
    except Exception as e:
        flash(f"Ошибка экспорта: {e}", "error")
        return redirect(url_for("index"))


@bp.route("/api/chapters/search")
def chapters_search():
    """Поиск по тексту всех глав."""
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    q = request.args.get("q", "").strip()
    if not q or len(q) < 2:
        return jsonify({"results": []})

    chapters = get_chapters(current["id"])
    results = []
    q_lower = q.lower()

    for ch in chapters:
        full = get_chapter(current["id"], ch["number"])
        if not full or not full.get("content"):
            continue
        content = full["content"]
        content_lower = content.lower()

        # Найти все вхождения с контекстом
        pos = 0
        matches = []
        while True:
            idx = content_lower.find(q_lower, pos)
            if idx == -1:
                break
            # Вырезать контекст: 100 символов до и после
            start = max(0, idx - 100)
            end = min(len(content), idx + len(q) + 100)
            snippet = content[start:end].replace("\n", " ")
            if start > 0:
                snippet = "..." + snippet
            if end < len(content):
                snippet = snippet + "..."
            matches.append(snippet)
            pos = idx + len(q)
            if len(matches) >= 3:  # максимум 3 совпадения на главу
                break

        if matches:
            results.append({
                "chapter_num": ch["number"],
                "chapter_title": ch.get("title") or f"Глава {ch['number']}",
                "matches": matches,
                "count": len(matches)
            })

    return jsonify({"results": results, "query": q, "total": len(results)})


@bp.route("/api/continuity/check", methods=["POST"])
def continuity_check():
    """Проверить главу на нарушения непрерывности."""
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    data = request.json or {}
    chapter_num = data.get("chapter_num")
    text = data.get("text", "").strip()
    # На странице главы выбора модели нет: проверяющий другой семьи, чем писатель
    # главы (plan_check.plan_checker_model — сначала «Настройки → Проверка
    # непрерывности»). Раньше брался из настройки как есть: после смены набора
    # ролей Sonnet проверял главы Sonnet (разбор 21.10)
    from engine.db import chapter_writers
    from engine.plan_check import plan_checker_model
    writer = chapter_writers(current["id"]).get(int(data.get("chapter_num") or 0), "")
    checker, why_none = plan_checker_model(writer) if writer else ("", "")
    model_value = checker if writer else page_model(data.get("model", ""), "continuity_model")
    if writer and not model_value:
        return jsonify({"error": why_none}), 400

    if not chapter_num or not text:
        return jsonify({"error": "Нужны chapter_num и text"}), 400
    if not model_value:
        return jsonify({"error": NO_MODEL}), 400

    try:
        from engine.continuity_checker import check_continuity_or_raise, skip_reason
        from engine.pipeline import call_llm as llm_call
        why = skip_reason(current["id"], int(chapter_num))
        if why:
            return jsonify({"ok": True, "violations": [], "count": 0, "skipped": why})

        def api_fn(prompt):
            return llm_call(model_value,
                           "Ты редактор серии. Ищешь нарушения непрерывности. Только JSON.",
                           prompt, max_tokens=600)

        violations = check_continuity_or_raise(current["id"], chapter_num, text, api_fn)
        return jsonify({"ok": True, "violations": violations, "count": len(violations)})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ─── L3 ──────────────────────────────────────────────────────────────────────

@bp.route("/api/l3/status")
def l3_status():
    current = get_current_project()
    if not current:
        return jsonify({"summaries": []})
    from engine.db import get_l3_status
    return jsonify({"summaries": get_l3_status(current["id"])})


@bp.route("/api/l3/<int:chapter_num>")
def l3_get(chapter_num):
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    from engine.db import get_l3_summary
    return jsonify({"summary": get_l3_summary(current["id"], chapter_num)})


@bp.route("/api/l3/<int:chapter_num>/generate", methods=["POST"])
def l3_generate(chapter_num):
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    ch = get_chapter(current["id"], chapter_num)
    if not ch:
        return jsonify({"error": "Глава не найдена"}), 404
    model_value = page_model((request.json or {}).get("model", ""))
    if not model_value:
        return jsonify({"error": NO_MODEL}), 400
    from engine.pipeline import generate_l3
    summary = generate_l3(current["id"], chapter_num, ch["content"],
                          get_cheap_model(model_value))
    if not summary:
        return jsonify({"error": "Не удалось сгенерировать саммари"}), 500
    return jsonify({"ok": True, "summary": summary})


@bp.route("/api/l3/<int:chapter_num>/delete", methods=["POST"])
def l3_delete(chapter_num):
    current = get_current_project()
    if not current:
        return jsonify({"error": "Нет проекта"}), 400
    from engine.db import delete_l3_summary
    delete_l3_summary(current["id"], chapter_num)
    return jsonify({"ok": True})

